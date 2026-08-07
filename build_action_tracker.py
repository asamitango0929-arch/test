# -*- coding: utf-8 -*-
"""社長向け：短時間勤務者採用プロジェクト アクション管理表を生成する。"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Meiryo"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "アクション管理"

# ---- 色・スタイル定義 -------------------------------------------------
navy = "1E2761"
ice = "CADCFC"
white = "FFFFFF"
grey = "F2F2F2"

thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

title_font = Font(name=FONT, size=16, bold=True, color=white)
sub_font = Font(name=FONT, size=9, color=white)
head_font = Font(name=FONT, size=10, bold=True, color=white)
cell_font = Font(name=FONT, size=10, color="222222")
cat_font = Font(name=FONT, size=10, bold=True, color=navy)

wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

# ---- タイトル行 -------------------------------------------------------
ws.merge_cells("A1:H1")
ws["A1"] = "短時間勤務者（パート・アルバイト）採用プロジェクト｜アクション管理表"
ws["A1"].font = title_font
ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws["A1"].fill = PatternFill("solid", fgColor=navy)
ws.row_dimensions[1].height = 30

ws.merge_cells("A2:H2")
ws["A2"] = "対象：社長／目的：説明会後に残る「検討中」論点を、担当・期日・ステータス付きで見える化し進捗管理する（更新日：2026-08-07）"
ws["A2"].font = sub_font
ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
ws["A2"].fill = PatternFill("solid", fgColor=navy)
ws.row_dimensions[2].height = 20

# ---- ヘッダー ---------------------------------------------------------
headers = ["No.", "カテゴリ", "Action（決めること）", "TODO（次にやること）",
           "担当", "期日", "区分", "ステータス"]
for c, h in enumerate(headers, start=1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.font = head_font
    cell.fill = PatternFill("solid", fgColor=navy)
    cell.alignment = center
    cell.border = border
ws.row_dimensions[4].height = 24

# ---- データ -----------------------------------------------------------
rows = [
    ("本社研修・教育", "短時間勤務者向けの本社研修の日数・カリキュラムを確定する",
     "フルタイム向け研修から短時間勤務者用に必要項目を切り出し、研修日数を設定", "教育担当", "26W稼働−8週", "検討中", "着手前"),
    ("交通費", "短時間勤務（週3日・1日4〜5h）に応じた交通費ルールを策定する",
     "日額上限／定期券支給要否／支給基準を人事規程に落とし込み決裁", "人事", "9月末", "検討中", "着手前"),
    ("試用期間", "試用期間の長さと評価・本採用フローを設定する",
     "試用期間の期間・評価項目・ライン投入判断のフローを整理し規程化", "人事", "9月末", "検討中", "着手前"),
    ("時給設定", "経験・資格に応じた時給テーブルを確定する（現状想定1,800〜2,100円）",
     "資格・経験区分ごとの時給レンジと昇給条件を確定、他基地展開の基準化", "人事", "8月末", "一部決定", "進行中"),
    ("雇用期間", "雇用契約期間（有期／更新条件）を設定する",
     "契約期間・更新基準・上限を法務確認のうえ雇用契約書ひな形に反映", "人事＋法務", "9月末", "検討中", "着手前"),
    ("資格付与・確認", "資格付与と既存資格の確認プロセスを整理する",
     "CKIN・GATE・ARR等、必要資格の付与手順と有資格者の棚卸し", "現場＋人事", "26W稼働−6週", "検討中", "着手前"),
    ("OJT・現場受入", "現場でのOJTと受入体制を整理する",
     "受入担当・OJT期間・チェックリストを支店ごとに整備", "現場＋教育", "26W稼働−4週", "検討中", "着手前"),
    ("シフト組込み", "短時間勤務者の勤務時間・シフトへの落とし込みルールを決める",
     "ピーク時間帯を軸に、既存シフトへ組み込む運用ルールを策定", "現場＋人事", "26W稼働−4週", "検討中", "着手前"),
    ("必要出面の確認", "Winter Schedule(26W)を基準に時間帯別の必要出面を確定する",
     "各支店で「何時に・どのポジションで・何人不足」を算出し集約", "各支店・所属長", "8月末", "検討中", "進行中"),
    ("採用人数の確定", "先行10名（成田・新千歳）を必要出面から逆算して再設定する",
     "不足出面を集約し、採用時間帯・人数を確定（10名は暫定値）", "人事", "9月中旬", "検討中", "進行中"),
    ("配置効果の検証", "「必要出面を何時間補完できたか」で効果を検証する仕組みを作る",
     "採用数ではなくライン投入・出面補完率でKPIを定義しモニタリング", "採用担当", "採用開始後", "検討中", "着手前"),
]

start = 5
for i, (cat, action, todo, owner, due, kind, status) in enumerate(rows):
    r = start + i
    fill = white if i % 2 == 0 else grey
    values = [i + 1, cat, action, todo, owner, due, kind, status]
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.border = border
        cell.fill = PatternFill("solid", fgColor=fill)
        if c == 1:
            cell.font = cell_font
            cell.alignment = center
        elif c == 2:
            cell.font = cat_font
            cell.alignment = wrap
        elif c in (5, 6, 7, 8):
            cell.font = cell_font
            cell.alignment = center
        else:
            cell.font = cell_font
            cell.alignment = wrap
    ws.row_dimensions[r].height = 42

# ---- 凡例 -------------------------------------------------------------
legend_r = start + len(rows) + 1
ws.merge_cells(start_row=legend_r, start_column=1, end_row=legend_r, end_column=8)
lc = ws.cell(row=legend_r, column=1,
             value="凡例）区分：説明会時点で『決定』か『検討中』か／ステータス：着手前・進行中・完了。"
                   "期日の「26W稼働−n週」はWinter Schedule稼働日からの逆算。担当・期日は説明会資料を基にした暫定案。")
lc.font = Font(name=FONT, size=8, italic=True, color="666666")
lc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.row_dimensions[legend_r].height = 28

# ---- 列幅・体裁 -------------------------------------------------------
widths = [5, 14, 34, 40, 13, 13, 9, 10]
for c, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(c)].width = w

ws.freeze_panes = "A5"
ws.sheet_view.showGridLines = False
ws.print_options.horizontalCentered = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

wb.save("短時間勤務者採用_アクション管理表.xlsx")
print("saved")
