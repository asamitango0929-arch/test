# -*- coding: utf-8 -*-
"""成田・朝の運用 日報 を生成する。

ポイント：
- 稼働便数・各種サマリーは「ハードコードの数値」ではなく実際の Excel 計算式
  （SUMPRODUCT / COUNTIF / MAX / INDEX-MATCH 等）で持たせ、便一覧を直せば
  自動で再計算されるようにしている。
- 日報の下部に Action / TODO / Responsible / STATUS / DUE DATE の欄を追加。

出典：週間線表・26W想定、成田空港事業部ハンドリング便（実線表データ）。
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import AreaChart, Reference
from openpyxl.utils import get_column_letter

FONT = "Meiryo"
navy, ice, amber, grey, white = "1E2761", "CADCFC", "E8A33D", "F2F2F2", "FFFFFF"
green, red, yellow = "2E7D32", "C0392B", "F6C445"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def F(sz, **kw):
    return Font(name=FONT, size=sz, **kw)

wrapL = Alignment(horizontal="left", vertical="center", wrap_text=True)
wrapTop = Alignment(horizontal="left", vertical="top", wrap_text=True)
ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
leftV = Alignment(horizontal="left", vertical="center")

REPORT_DATE = "2026-08-07"   # 日報の対象日（更新日）

def hm(s):
    h, m = map(int, s.split(":"))
    return h * 60 + m

# ---- 実データ（線表より）: (carrier, 便, 行先, STA, STD, CTR OPEN, 区分, 運航日, 配置) ----
flights = [
    ("AM", "AM058/057", "MEX", "06:30", "09:35", "06:30", "朝", "毎日", "◯"),
    ("VN", "VN318/319", "DAD", "07:35", "09:00", "06:00", "朝", "毎日", "◯"),
    ("VN", "VN310/311", "HAN", "07:35", "09:30", "06:30", "朝", "毎日", "◯"),
    ("VN", "VN306/307", "SGN", "08:00", "09:30", "06:30", "朝", "毎日", "◯"),
    ("VN", "VN316/317", "DAD", "08:30", "10:30", "07:30", "朝", "月木金土日", "◯"),
    ("VJ", "VJ822/823", "SGN", "07:40", "08:55", "05:55", "朝", "毎日", "◯"),
    ("VJ", "VJ932/933", "HAN", "08:00", "09:30", "06:30", "朝", "毎日", "◯"),
    ("5J", "5J5062/063", "CEB", "08:10", "08:55", "05:55", "朝", "毎日", "◯"),
    ("5J", "5J5068/069", "CRK", "10:25", "11:15", "08:15", "朝", "毎日", "◯"),
    ("RF", "RF392/391", "CJJ", "09:40", "10:40", "08:10", "朝", "毎日", "◯"),
    ("UL", "UL454/455", "CMB", "08:10", "11:15", "08:15", "朝", "火木土", "◯"),
    ("5J", "5J5054/055", "MNL", "12:25", "13:45", "10:45", "昼", "毎日", "△(昼)"),
    ("VJ", "VJ934/935", "HAN", "15:30", "16:30", "13:30", "夕", "毎日", "△(Winter条件付)"),
    ("5J", "5J5056/057", "MNL", "18:00", "19:15", "16:15", "夕", "毎日", "△(夕)"),
]

# ---- 下部：Action / TODO / Responsible / STATUS / DUE DATE ----
# (Action, TODO, Responsible, STATUS, DUE DATE)
actions = [
    ("Winter Schedule(26W)を基準に時間帯別の必要出面を確定する",
     "各支店で『何時に・どのポジションで・何人不足』を算出し集約",
     "各支店・所属長", "進行中", "8月末"),
    ("先行10名（成田・新千歳）を必要出面から逆算して再設定する",
     "不足出面を集約し、採用時間帯・人数を確定（10名は暫定値）",
     "人事", "進行中", "9月中旬"),
    ("短時間勤務者の勤務時間・シフトへの落とし込みルールを決める",
     "朝ピーク（06:30〜09:30）を軸に既存シフトへ組み込む運用ルールを策定",
     "現場＋人事", "着手前", "26W稼働−4週"),
    ("配置候補便（◯）への有資格者の割当てを整理する",
     "CKIN・GATE・ARR等、必要資格の付与手順と有資格者の棚卸し",
     "現場＋人事", "着手前", "26W稼働−6週"),
    ("『必要出面を何時間補完できたか』で効果を検証する仕組みを作る",
     "採用数ではなくライン投入・出面補完率でKPIを定義しモニタリング",
     "採用担当", "着手前", "採用開始後"),
]

status_color = {"完了": green, "進行中": yellow, "着手前": red}

wb = openpyxl.Workbook()

# =====================================================================
# Sheet1: 便一覧（ソースデータ）── 計算式のもとになる隠しヘルパー列つき
# =====================================================================
ws1 = wb.active
ws1.title = "便一覧"
ws1.sheet_view.showGridLines = False

ws1.merge_cells("A1:L1")
ws1["A1"] = "成田・地上ハンドリング便 一覧（実線表ベース・計算式のソースデータ）"
ws1["A1"].font = F(14, bold=True, color=white)
ws1["A1"].fill = PatternFill("solid", fgColor=navy)
ws1["A1"].alignment = leftV
ws1.row_dimensions[1].height = 26

heads1 = ["No.", "Carrier", "便名", "行先", "到着 STA", "出発 STD", "CTR OPEN",
          "時間帯", "運航日", "パート配置", "開始(分)", "終了(分)"]
for c, h in enumerate(heads1, 1):
    cell = ws1.cell(3, c, h)
    cell.font = F(10, bold=True, color=white)
    cell.fill = PatternFill("solid", fgColor=navy)
    cell.alignment = ctr
    cell.border = border
ws1.row_dimensions[3].height = 22

DATA0 = 4  # 最初のデータ行
for i, (car, flt, dest, sta, std, cto, band, days, place) in enumerate(flights):
    r = DATA0 + i
    base = white if band == "朝" else "FBEAD2"
    vals = [i + 1, car, flt, dest, sta, std, cto, band, days, place]
    for c, v in enumerate(vals, 1):
        cell = ws1.cell(r, c, v)
        cell.border = border
        cell.fill = PatternFill("solid", fgColor=base)
        cell.font = F(10, bold=(c == 2), color="222222")
        cell.alignment = ctr if c != 3 else wrapL
    # ヘルパー列（計算式）: 開始(分)=MIN(CTR OPEN, STA), 終了(分)=STD
    # 時刻テキストを分に変換：HOUR/MINUTE ではなく TIMEVALUE を利用
    ws1.cell(r, 11, f'=MIN(TIMEVALUE(G{r}),TIMEVALUE(E{r}))*1440').font = F(9, color="888888")
    ws1.cell(r, 12, f'=TIMEVALUE(F{r})*1440').font = F(9, color="888888")
    ws1.cell(r, 11).border = border
    ws1.cell(r, 12).border = border
    ws1.cell(r, 11).alignment = ctr
    ws1.cell(r, 12).alignment = ctr
    ws1.row_dimensions[r].height = 20
DATAN = DATA0 + len(flights) - 1  # 最終データ行

for c, w in enumerate([5, 9, 12, 8, 10, 10, 10, 8, 12, 16, 9, 9], 1):
    ws1.column_dimensions[get_column_letter(c)].width = w
# ヘルパー列は非表示
ws1.column_dimensions["K"].hidden = True
ws1.column_dimensions["L"].hidden = True
ws1.freeze_panes = "A4"

note1 = DATAN + 2
ws1.merge_cells(start_row=note1, start_column=1, end_row=note1, end_column=12)
ws1.cell(note1, 1,
         "配置：◯=朝ピークの配置候補（VN/AM/RF/5J/UL/VJ）／△=夕・昼便。"
         "K・L列（開始/終了・分）は稼働便数の計算式が参照するヘルパー列（非表示）。数値はすべて線表の実データ。")
ws1.cell(note1, 1).font = F(8, italic=True, color="666666")
ws1.cell(note1, 1).alignment = wrapL
ws1.row_dimensions[note1].height = 28

# =====================================================================
# Sheet2: 時間帯別 稼働便数 ── SUMPRODUCT の計算式で自動集計
# =====================================================================
ws2 = wb.create_sheet("時間帯別")
ws2.sheet_view.showGridLines = False
ws2.merge_cells("A1:F1")
ws2["A1"] = "成田・時間帯別 稼働便数（30分刻み・計算式で自動集計）"
ws2["A1"].font = F(14, bold=True, color=white)
ws2["A1"].fill = PatternFill("solid", fgColor=navy)
ws2["A1"].alignment = leftV
ws2.row_dimensions[1].height = 26
ws2.merge_cells("A2:F2")
ws2["A2"] = "稼働便数 = 便一覧の[開始(分)〜終了(分)]がその時刻を含む便数。時刻を直せば自動再計算される。"
ws2["A2"].font = F(9, color=white)
ws2["A2"].fill = PatternFill("solid", fgColor=navy)
ws2["A2"].alignment = leftV
ws2.row_dimensions[2].height = 18

for c, h in enumerate(["時刻", "稼働便数", "（分）"], 1):
    cell = ws2.cell(4, c, h)
    cell.font = F(10, bold=True, color=white)
    cell.fill = PatternFill("solid", fgColor=navy)
    cell.alignment = ctr
    cell.border = border
ws2.row_dimensions[4].height = 20

t0, t1, step = hm("05:30"), hm("20:00"), 30
SROW = 5
r = SROW
t = t0
while t < t1:
    lab = f"{t//60:02d}:{t%60:02d}"
    a = ws2.cell(r, 1, lab)
    a.font = F(9)
    a.alignment = ctr
    a.border = border
    # C列に分（計算式が使う）
    cmin = ws2.cell(r, 3, t)
    cmin.font = F(8, color="AAAAAA")
    cmin.alignment = ctr
    cmin.border = border
    # B列：SUMPRODUCT で同時稼働便数
    b = ws2.cell(
        r, 2,
        f"=SUMPRODUCT((便一覧!$K${DATA0}:$K${DATAN}<=C{r})*"
        f"(便一覧!$L${DATA0}:$L${DATAN}>C{r}))")
    b.font = F(9)
    b.alignment = ctr
    b.border = border
    is_morning = t0 <= t < hm("10:00")
    b.fill = PatternFill("solid", fgColor=ice if is_morning else white)
    r += 1
    t += step
SLAST = r - 1
ws2.column_dimensions["C"].hidden = True

# 合計・ピーク（計算式）
sr = SLAST + 2
ws2.cell(sr, 1, "延べ稼働（合計）").font = F(9, bold=True, color=navy)
ws2.cell(sr, 2, f"=SUM(B{SROW}:B{SLAST})").font = F(9, bold=True)
ws2.cell(sr + 1, 1, "ピーク同時稼働").font = F(9, bold=True, color=navy)
ws2.cell(sr + 1, 2, f"=MAX(B{SROW}:B{SLAST})").font = F(9, bold=True)
ws2.cell(sr + 2, 1, "ピーク時刻").font = F(9, bold=True, color=navy)
ws2.cell(sr + 2, 2,
         f"=INDEX(A{SROW}:A{SLAST},MATCH(MAX(B{SROW}:B{SLAST}),B{SROW}:B{SLAST},0))").font = F(9, bold=True)
for rr in (sr, sr + 1, sr + 2):
    ws2.cell(rr, 1).alignment = leftV
    ws2.cell(rr, 2).alignment = ctr

chart = AreaChart()
chart.title = "時間帯別 稼働便数（計算式ベース）"
chart.style = 2
chart.height = 9
chart.width = 20
data = Reference(ws2, min_col=2, min_row=4, max_row=SLAST)
cats = Reference(ws2, min_col=1, min_row=SROW, max_row=SLAST)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.y_axis.title = "同時稼働便数"
chart.x_axis.title = "時刻"
chart.legend = None
s = chart.series[0]
s.graphicalProperties.solidFill = navy
s.graphicalProperties.line.solidFill = navy
ws2.add_chart(chart, "E4")

ws2.column_dimensions["A"].width = 16
ws2.column_dimensions["B"].width = 11
ws2.freeze_panes = "A5"

# =====================================================================
# Sheet3: 日報（サマリー計算式 ＋ 下部に Action/TODO/Responsible/STATUS/DUE DATE）
# =====================================================================
ws = wb.create_sheet("日報")
ws.sheet_view.showGridLines = False

ws.merge_cells("A1:E1")
ws["A1"] = "成田・朝の運用 日報"
ws["A1"].font = F(18, bold=True, color=white)
ws["A1"].fill = PatternFill("solid", fgColor=navy)
ws["A1"].alignment = leftV
ws.row_dimensions[1].height = 34
ws.merge_cells("A2:E2")
ws["A2"] = f"対象：成田空港事業部ハンドリング便（26W想定）／更新日：{REPORT_DATE}／数値はすべて計算式で自動集計"
ws["A2"].font = F(9, color=white)
ws["A2"].fill = PatternFill("solid", fgColor=navy)
ws["A2"].alignment = leftV
ws.row_dimensions[2].height = 18

# ---- 運用サマリー（KPI・すべて計算式） ----
ws.merge_cells("A4:E4")
ws["A4"] = "■ 運用サマリー（KPI）"
ws["A4"].font = F(12, bold=True, color=navy)
ws.row_dimensions[4].height = 22

# (ラベル, 計算式, 単位/書式)
kpis = [
    ("総取扱便数", f"=COUNTA(便一覧!$C${DATA0}:$C${DATAN})", "便"),
    ("朝便数（06:00〜10:59）", f'=COUNTIF(便一覧!$H${DATA0}:$H${DATAN},"朝")', "便"),
    ("朝便の比率", f'=COUNTIF(便一覧!$H${DATA0}:$H${DATAN},"朝")/COUNTA(便一覧!$C${DATA0}:$C${DATAN})', "%"),
    ("パート配置候補（◯）", f'=COUNTIF(便一覧!$J${DATA0}:$J${DATAN},"◯")', "便"),
    ("ピーク同時稼働便数", f"=MAX(時間帯別!$B${SROW}:$B${SLAST})", "便"),
    ("ピーク時刻", f"=INDEX(時間帯別!$A${SROW}:$A${SLAST},MATCH(MAX(時間帯別!$B${SROW}:$B${SLAST}),時間帯別!$B${SROW}:$B${SLAST},0))", "時刻"),
    ("延べ稼働（30分×便）", f"=SUM(時間帯別!$B${SROW}:$B${SLAST})", "コマ"),
    ("朝(〜10:00)延べ稼働の比率",
     f"=SUMPRODUCT((時間帯別!$C${SROW}:$C${SLAST}<{hm('10:00')})*時間帯別!$B${SROW}:$B${SLAST})/SUM(時間帯別!$B${SROW}:$B${SLAST})",
     "%"),
]
kr = 5
for label, formula, unit in kpis:
    lc = ws.cell(kr, 1, label)
    lc.font = F(10, color="222222")
    lc.alignment = wrapL
    lc.fill = PatternFill("solid", fgColor=grey)
    lc.border = border
    ws.merge_cells(start_row=kr, start_column=1, end_row=kr, end_column=3)
    for cc in range(1, 4):
        ws.cell(kr, cc).border = border
        ws.cell(kr, cc).fill = PatternFill("solid", fgColor=grey)
    vc = ws.cell(kr, 4, formula)
    vc.font = F(12, bold=True, color=navy)
    vc.alignment = ctr
    vc.border = border
    if unit == "%":
        vc.number_format = "0.0%"
    uc = ws.cell(kr, 5, unit if unit != "%" else "")
    uc.font = F(9, color="666666")
    uc.alignment = Alignment(horizontal="left", vertical="center")
    uc.border = border
    ws.row_dimensions[kr].height = 22
    kr += 1

# ---- 下部：Action / TODO / Responsible / STATUS / DUE DATE ----
ar = kr + 1
ws.merge_cells(start_row=ar, start_column=1, end_row=ar, end_column=5)
ws.cell(ar, 1, "■ アクション（Action / TODO / Responsible / STATUS / DUE DATE）")
ws.cell(ar, 1).font = F(12, bold=True, color=navy)
ws.row_dimensions[ar].height = 22
ar += 1

ahead = ["Action", "TODO", "Responsible", "STATUS", "DUE DATE"]
for c, h in enumerate(ahead, 1):
    cell = ws.cell(ar, c, h)
    cell.font = F(10, bold=True, color=white)
    cell.fill = PatternFill("solid", fgColor=navy)
    cell.alignment = ctr
    cell.border = border
ws.row_dimensions[ar].height = 22

astart = ar + 1
for i, (action, todo, resp, status, due) in enumerate(actions):
    r = astart + i
    fill = white if i % 2 == 0 else grey
    vals = [action, todo, resp, status, due]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(r, c, v)
        cell.border = border
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = F(10, color="222222")
        if c in (1, 2):
            cell.alignment = wrapTop
        else:
            cell.alignment = ctr
    # STATUS セルに色を付ける
    sc = ws.cell(r, 4)
    if status in status_color:
        sc.font = F(10, bold=True, color=white)
        sc.fill = PatternFill("solid", fgColor=status_color[status])
    ws.row_dimensions[r].height = 40
aend = astart + len(actions) - 1

# 凡例
lr = aend + 1
ws.merge_cells(start_row=lr, start_column=1, end_row=lr, end_column=5)
ws.cell(lr, 1,
        "凡例）STATUS：着手前（赤）／進行中（黄）／完了（緑）。DUE DATE の「26W稼働−n週」は"
        "Winter Schedule 稼働日からの逆算。KPI・稼働便数はすべて計算式（便一覧を直せば自動再計算）。")
ws.cell(lr, 1).font = F(8, italic=True, color="666666")
ws.cell(lr, 1).alignment = wrapL
ws.row_dimensions[lr].height = 30

for c, w in enumerate([40, 44, 14, 12, 16], 1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = "A5"
ws.print_options.horizontalCentered = True
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# 日報シートを先頭に
wb.move_sheet("日報", -(len(wb.sheetnames) - 1))
wb.active = wb.sheetnames.index("日報")

# 開いた時に計算式を強制再計算（キャッシュ値がなくても値が表示される）
wb.calculation.fullCalcOnLoad = True

wb.save("成田_朝の運用_日報.xlsx")
print("saved: 成田_朝の運用_日報.xlsx  flights rows", DATA0, "-", DATAN, " slot rows", SROW, "-", SLAST)
