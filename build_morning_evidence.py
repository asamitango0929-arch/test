# -*- coding: utf-8 -*-
"""成田・朝集中の客観的証拠＋短時間勤務者 配置候補便 を実線表データから生成。
出典：週間線表(6ea4fb69-___9Gxlsx.xlsx) MON〜SUN、成田空港事業部ハンドリング便。"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import AreaChart, Reference
from openpyxl.utils import get_column_letter

FONT = "Meiryo"
navy, ice, amber, grey, white = "1E2761", "CADCFC", "E8A33D", "F2F2F2", "FFFFFF"
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def F(sz, **kw): return Font(name=FONT, size=sz, **kw)
wrapL = Alignment(horizontal="left", vertical="center", wrap_text=True)
ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)

# ---- 実データ（線表より）: (carrier, 便, 行先, STA, STD, CTR OPEN, 区分, 運航日, 配置) ----
def hm(s): h, m = map(int, s.split(":")); return h*60+m
flights = [
    ("AM", "AM058/057", "MEX", "06:30", "09:35", "06:30", "朝", "毎日", "◯"),
    ("VN", "VN318/319", "DAD", "07:35", "09:00", "06:00", "朝", "毎日", "◯"),
    ("VN", "VN310/311", "HAN", "07:35", "09:30", "06:30", "朝", "毎日", "◯"),
    ("VN", "VN306/307", "SGN", "08:00", "09:30", "06:30", "朝", "毎日", "◯"),
    ("VN", "VN316/317", "DAD", "08:30", "10:30", "07:30", "朝", "月木金土日", "◯"),
    ("VJ", "VJ822/823", "SGN", "07:40", "08:55", "05:55", "朝", "毎日", "◯"),
    ("VJ", "VJ932/933", "HAN", "08:00", "09:30", "06:30", "朝", "毎日", "◯"),
    ("5J", "5J5062/063", "CEB", "08:10", "08:55", "05:55", "朝", "毎日", "◯"),
    ("5J", "5J5068/069", "CRK", "10:25", "11:15", "08:15", "朝", "毎日", "◯"),
    ("RF", "RF392/391", "CJJ", "09:40", "10:40", "08:10", "朝", "毎日", "◯"),
    ("UL", "UL454/455", "CMB", "08:10", "11:15", "08:15", "朝", "火木土", "◯"),
    ("5J", "5J5054/055", "MNL", "12:25", "13:45", "10:45", "昼", "毎日", "△(昼)"),
    ("VJ", "VJ934/935", "HAN", "15:30", "16:30", "13:30", "夕", "毎日", "△(Winter条件付)"),
    ("5J", "5J5056/057", "MNL", "18:00", "19:15", "16:15", "夕", "毎日", "△(夕)"),
]

wb = openpyxl.Workbook()

# =====================================================================
# Sheet1: 配置候補便一覧
# =====================================================================
ws = wb.active
ws.title = "配置候補便_線表"
ws.sheet_view.showGridLines = False

ws.merge_cells("A1:J1")
ws["A1"] = "成田・短時間勤務者 配置候補便一覧（実線表ベース）"
ws["A1"].font = F(15, bold=True, color=white); ws["A1"].fill = PatternFill("solid", fgColor=navy)
ws["A1"].alignment = Alignment(horizontal="left", vertical="center"); ws.row_dimensions[1].height = 28
ws.merge_cells("A2:J2")
ws["A2"] = "対象：成田空港事業部がハンドリングを担当する便／出典：週間線表・26W想定。時刻はSTA=到着・STD=出発・CTR OPEN=カウンター開設。"
ws["A2"].font = F(9, color=white); ws["A2"].fill = PatternFill("solid", fgColor=navy)
ws["A2"].alignment = Alignment(horizontal="left", vertical="center"); ws.row_dimensions[2].height = 18

heads = ["No.", "Carrier", "便名", "行先", "到着 STA", "出発 STD", "CTR OPEN", "時間帯", "運航日", "パート配置"]
for c, h in enumerate(heads, 1):
    cell = ws.cell(4, c, h); cell.font = F(10, bold=True, color=white)
    cell.fill = PatternFill("solid", fgColor=navy); cell.alignment = ctr; cell.border = border
ws.row_dimensions[4].height = 22

for i, (car, flt, dest, sta, std, cto, band, days, place) in enumerate(flights):
    r = 5 + i
    base = white if band == "朝" else "FBEAD2"  # 夕/昼はうっすらアンバー
    vals = [i+1, car, flt, dest, sta, std, cto, band, days, place]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(r, c, v); cell.border = border
        cell.fill = PatternFill("solid", fgColor=base)
        cell.font = F(10, bold=(c == 2), color="222222")
        cell.alignment = ctr if c != 3 else wrapL
    ws.row_dimensions[r].height = 20

note_r = 5 + len(flights) + 1
ws.merge_cells(start_row=note_r, start_column=1, end_row=note_r, end_column=10)
ws.cell(note_r, 1, "配置：◯=朝ピークの配置候補（VN/AM/RF/5J/UL/VJ）／△=夕・昼便。VJ934/935（夕方HAN）はWinterで運休の場合は対象外、運航する場合は別途候補。"
                  "ULは火・木・土の運航。VN316/317は一部曜日運休。数値はすべて線表の実データ。")
ws.cell(note_r, 1).font = F(8, italic=True, color="666666"); ws.cell(note_r, 1).alignment = wrapL
ws.row_dimensions[note_r].height = 30
for c, w in enumerate([5, 9, 12, 8, 10, 10, 10, 8, 11, 15], 1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.freeze_panes = "A5"

# =====================================================================
# Sheet2: 時間帯別 稼働便数（客観的証拠の折れ線/面グラフ）
# =====================================================================
ws2 = wb.create_sheet("時間帯別_稼働便数")
ws2.sheet_view.showGridLines = False
ws2.merge_cells("A1:F1")
ws2["A1"] = "成田・時間帯別 地上ハンドリング稼働便数（実線表ベース）"
ws2["A1"].font = F(15, bold=True, color=white); ws2["A1"].fill = PatternFill("solid", fgColor=navy)
ws2["A1"].alignment = Alignment(horizontal="left", vertical="center"); ws2.row_dimensions[1].height = 28
ws2.merge_cells("A2:F2")
ws2["A2"] = "各便の地上作業時間帯［CTR OPEN〜出発STD、到着便は到着STAまで含む］を30分刻みで集計。朝に集中し、昼以降に急減することが分かる。"
ws2["A2"].font = F(9, color=white); ws2["A2"].fill = PatternFill("solid", fgColor=navy)
ws2["A2"].alignment = Alignment(horizontal="left", vertical="center"); ws2.row_dimensions[2].height = 18

# 稼働ウィンドウ = [min(CTR OPEN, STA), STD]
windows = []
for (car, flt, dest, sta, std, cto, band, days, place) in flights:
    start = min(hm(cto), hm(sta)); end = hm(std)
    windows.append((start, end))

# 05:30〜20:00 を30分刻み
t0, t1, step = hm("05:30"), hm("20:00"), 30
ws2.cell(4, 1, "時刻").font = F(10, bold=True, color=white)
ws2.cell(4, 2, "稼働便数").font = F(10, bold=True, color=white)
for c in (1, 2):
    ws2.cell(4, c).fill = PatternFill("solid", fgColor=navy); ws2.cell(4, c).alignment = ctr; ws2.cell(4, c).border = border
row = 5
t = t0
while t < t1:
    cnt = sum(1 for (s, e) in windows if s <= t < e)
    lab = f"{t//60:02d}:{t % 60:02d}"
    a = ws2.cell(row, 1, lab); a.font = F(9); a.alignment = ctr; a.border = border
    b = ws2.cell(row, 2, cnt); b.font = F(9); b.alignment = ctr; b.border = border
    b.fill = PatternFill("solid", fgColor=ice if 5 <= (t-t0)//30 else white)
    row += 1; t += step
last = row - 1

chart = AreaChart()
chart.title = "時間帯別 稼働便数（成田空港事業部ハンドリング便）"
chart.style = 2
chart.height = 9; chart.width = 22
data = Reference(ws2, min_col=2, min_row=4, max_row=last)
cats = Reference(ws2, min_col=1, min_row=5, max_row=last)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.y_axis.title = "同時稼働便数"
chart.x_axis.title = "時刻"
chart.legend = None
s = chart.series[0]
from openpyxl.chart.marker import DataPoint
from openpyxl.drawing.fill import PatternFillProperties, ColorChoice
s.graphicalProperties.solidFill = navy
s.graphicalProperties.line.solidFill = navy
ws2.add_chart(chart, "D4")

ws2.column_dimensions["A"].width = 8
ws2.column_dimensions["B"].width = 10
nr = last + 2
ws2.merge_cells(start_row=nr, start_column=1, end_row=nr, end_column=2)
ws2.cell(nr, 1, "※ UL・VN316等の特定曜日便も含む代表パターン。夕方の小さな山はVJ934/935・5J5056/057。").font = F(8, italic=True, color="666666")
ws2.cell(nr, 1).alignment = wrapL; ws2.row_dimensions[nr].height = 26

wb.save("成田_朝集中_分析.xlsx")
print("saved rows=", last)
