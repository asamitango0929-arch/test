# -*- coding: utf-8 -*-
"""数式のキャッシュ値を評価してxlsxに書き込む（LibreOffice不在環境向け）。

数式そのものはXML上に残すため、Excelで開けば通常どおり再計算される。
"""
import re, sys, zipfile, shutil
from xml.etree import ElementTree as ET
import openpyxl
from openpyxl.utils import range_boundaries, get_column_letter

PATH = sys.argv[1]
wb = openpyxl.load_workbook(PATH)

REF = re.compile(r"(?:'([^']+)'|([A-Za-z_぀-鿿][\w぀-鿿]*))!")
CELL = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
cache = {}

def cell_value(sheet, coord):
    key = (sheet, coord.replace("$", ""))
    if key in cache:
        v = cache[key]
        if v is _PENDING:
            raise RuntimeError(f"循環参照: {key}")
        return v
    cache[key] = _PENDING
    raw = wb[sheet][key[1]].value
    if isinstance(raw, str) and raw.startswith("="):
        val = evaluate(raw[1:], sheet)
    elif isinstance(raw, (int, float)):
        val = raw
    else:
        val = 0
    cache[key] = val
    return val

class _P: pass
_PENDING = _P()

def expand(rng, sheet):
    mn_c, mn_r, mx_c, mx_r = range_boundaries(rng.replace("$", ""))
    return [cell_value(sheet, f"{get_column_letter(c)}{r}")
            for r in range(mn_r, mx_r + 1) for c in range(mn_c, mx_c + 1)]

def evaluate(expr, sheet):
    # SUM / AVERAGE を先に畳み込む
    def fold(m):
        fn, arg = m.group(1).upper(), m.group(2)
        sh, rng = sheet, arg
        r = REF.match(arg)
        if r:
            sh = r.group(1) or r.group(2)
            rng = arg[r.end():]
        vals = expand(rng, sh)
        return repr(sum(vals) if fn == "SUM" else sum(vals) / len(vals))
    prev = None
    while prev != expr:
        prev = expr
        expr = re.sub(r"\b(SUM|AVERAGE)\(([^()]+)\)", fold, expr, flags=re.I)
    # 単独セル参照を値に置換（シート指定あり／なし）
    def sub_ref(m):
        sh = m.group(1) or m.group(2)
        return f"\x00{sh}\x00"
    expr = REF.sub(sub_ref, expr)
    def sub_cell(m):
        start = m.start()
        pre = expr[:start]
        sh = sheet
        tail = re.search(r"\x00([^\x00]+)\x00$", pre)
        if tail:
            sh = tail.group(1)
        return repr(cell_value(sh, m.group(0)))
    expr = CELL.sub(sub_cell, expr)
    expr = re.sub(r"\x00[^\x00]+\x00", "", expr)
    return eval(expr, {"__builtins__": {}}, {})

# 全数式セルを評価
targets = {}
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                targets.setdefault(ws.title, {})[c.coordinate] = cell_value(ws.title, c.coordinate)

sheet_file = {ws.title: f"xl/worksheets/sheet{i+1}.xml" for i, ws in enumerate(wb.worksheets)}
for t, vals in targets.items():
    print(f"[{t}] {len(vals)} 数式セルを評価")

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ET.register_namespace("", NS)

tmp = PATH + ".tmp"
with zipfile.ZipFile(PATH) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        match = [t for t, f in sheet_file.items() if f == item.filename]
        if match and match[0] in targets:
            vals = targets[match[0]]
            root = ET.fromstring(data)
            n = 0
            for c in root.iter(f"{{{NS}}}c"):
                f = c.find(f"{{{NS}}}f")
                if f is None:
                    continue
                ref = c.get("r")
                if ref not in vals:
                    continue
                for old in c.findall(f"{{{NS}}}v"):
                    c.remove(old)
                v = ET.SubElement(c, f"{{{NS}}}v")
                val = vals[ref]
                v.text = repr(round(val, 10)) if isinstance(val, float) else str(val)
                c.attrib.pop("t", None)
                n += 1
            print(f"  -> {item.filename}: {n} セルに値を書き込み")
            data = ET.tostring(root, encoding="UTF-8", xml_declaration=True)
        zout.writestr(item, data)
shutil.move(tmp, PATH)
print("完了")
