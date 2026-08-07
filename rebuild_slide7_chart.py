# -*- coding: utf-8 -*-
"""slide7の拠点別グラフを python-pptx でネイティブ再構築（replace_data由来の非表示を解消）。
元と同じ緑・データラベル付き横棒、値は万円。"""
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

prs = Presentation("out_deck.pptx")
s7 = prs.slides[6]
old = [sh for sh in s7.shapes if sh.has_chart][0]
L, T, W, H = old.left, old.top, old.width, old.height
old._element.getparent().remove(old._element)   # 壊れたチャートを除去

cd = CategoryChartData()
cd.categories = ["那覇", "羽田", "千歳", "北九州", "関空", "成田"]
cd.add_series("旅客 営業損益", [760, 330, 130, 0, -770, -1730])   # 万円/月
gf = s7.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, L, T, W, H, cd)
ch = gf.chart
ch.has_legend = False
ch.has_title = False
plot = ch.plots[0]
plot.gap_width = 150
plot.vary_by_categories = False
ser = ch.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = RGBColor.from_string("1E7A5A")
plot.has_data_labels = True
dl = plot.data_labels
dl.number_format = "#,##0"
dl.number_format_is_linked = False
dl.font.size = Pt(11); dl.font.name = "Meiryo"; dl.font.bold = True
dl.font.color.rgb = RGBColor.from_string("1F2733")
for ax in (ch.category_axis, ch.value_axis):
    ax.tick_labels.font.size = Pt(10)
    ax.tick_labels.font.name = "Meiryo"
    ax.tick_labels.font.color.rgb = RGBColor.from_string("6B7686")

prs.save("out_deck.pptx")
print("rebuilt slide7 chart")
