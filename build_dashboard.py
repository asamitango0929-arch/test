# -*- coding: utf-8 -*-
"""短時間勤務者 採用KPIダッシュボード＋社長報告フォーマット ビルダー"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.workbook.properties import CalcProperties

OUT = "/home/user/test/短時間勤務者_採用KPIダッシュボード.xlsx"
JP = "Meiryo"

# ---------- palette ----------
NAVY="1F3864"; BLUE="2E5496"; BLUE_L="D9E1F2"; GREY_L="F2F2F2"; GREY_H="808080"
WHITE="FFFFFF"; INPUT="FFF2CC"; INK="1A1A1A"; GOLD="7F6000"
GREEN_F="C6EFCE"; GREEN_T="006100"; AMBER_F="FFEB9C"; AMBER_T="9C6500"; RED_F="FFC7CE"; RED_T="9C0006"

thin=Side(style="thin",color="BFBFBF")
BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
def fill(h): return PatternFill("solid",fgColor=h)
def font(sz=11,b=False,color=INK,it=False): return Font(name=JP,size=sz,bold=b,color=color,italic=it)
CEN=Alignment(horizontal="center",vertical="center",wrap_text=True)
LEF=Alignment(horizontal="left",vertical="center",wrap_text=True)

A="応募者データ"; DATA_END=502; DL_LAST=200
# applicant column map
COL={'no':'A','媒体':'B','基地':'C','応募日':'D','氏名':'E','年齢':'F','性別':'G','国籍':'H',
     '経験':'I','属性':'J','書類選考':'K','書類NG':'L','面接':'M','面接NG':'N','採用':'O','備考':'P'}
def rng(k): return f"{A}!${COL[k]}$3:${COL[k]}${DATA_END}"

# funnel live formulas (accept optional extra criteria for base/media)
def f_ouboe(extra=""):   return f'=COUNTA({rng("氏名")})' if not extra else f'=COUNTIFS({extra})'
def countif(k,v):        return f'COUNTIF({rng(k)},"{v}")'
def countifs(pairs):     return "COUNTIFS(" + ",".join(f'{rng(k)},"{v}"' for k,v in pairs) + ")"

wb=openpyxl.Workbook()

# =====================================================================
#  MASTER
# =====================================================================
mst=wb.active; mst.title="マスタ"
master={
 "媒体":["バイトル","マイナビバイト","ちゃんと","社員紹介","HP","その他"],
 "基地":["NRT","CTS"],
 "性別":["男","女"],
 "国籍":["日本","外国籍"],
 "経験":["経験者","未経験者"],
 "属性":["学生","フリーター","社会人","主婦・主夫","その他"],
 "書類選考":["未対応","確認中","通過","保留","不通過","辞退","連絡不通"],
 "書類NG理由":["年齢","通勤距離","国籍","日本語力","英語力","勤務時間不一致","その他"],
 "面接":["未設定","面接調整中","面接日確定","面接実施済み"],
 "面接NG理由":["勤務条件不一致","シフト条件不一致","経験不足","接客適性","コミュニケーション","日本語力",
             "英語力","航空保安上の要件","通勤困難","給与条件不一致","入社可能時期が合わない","他社内定",
             "本人辞退","総合判断","その他"],
 "採用":["選考中","採用","不採用","採用辞退","入社予定","入社辞退","保留"],
}
mst["A1"]="マスタ（プルダウン用リスト）"; mst["A1"].font=font(13,True,BLUE)
mrange={}
for c,(head,vals) in enumerate(master.items(),1):
    col=get_column_letter(c)
    h=mst.cell(row=2,column=c,value=head); h.font=font(11,True,WHITE); h.fill=fill(BLUE); h.alignment=CEN; h.border=BORDER
    for r,v in enumerate(vals,3):
        cel=mst.cell(row=r,column=c,value=v); cel.font=font(10); cel.border=BORDER; cel.alignment=LEF
    mst.column_dimensions[col].width=16 if head in("面接NG理由","書類NG理由") else 12
    mrange[head]=(col,2+len(vals))
mst.freeze_panes="A3"

# =====================================================================
#  APPLICANT DATA
# =====================================================================
app=wb.create_sheet("応募者データ")
cols=["No.","媒体","基地","応募日","氏名","年齢","性別","国籍","経験","属性",
      "書類選考","書類NG理由","面接","面接NG理由","採用","備考"]
widths=[5,13,7,11,16,6,6,8,9,10,10,12,12,14,10,26]
# base rows (経験/属性 は原則未入力＝担当者記入。事実が明確な高校生のみ 学生)
data=[
 [1,"バイトル","NRT",46227,"日高 秀光",73,"男","日本","","","不通過","年齢","未設定","","不採用",""],
 [2,"バイトル","NRT",46230,"浅利 真弓",63,"女","日本","","","不通過","年齢","未設定","","不採用",""],
 [3,"バイトル","NRT",46231,"浅木 洋二",76,"男","日本","","","不通過","年齢","未設定","","不採用",""],
 [4,"バイトル","NRT",46231,"佐伯 幸男",55,"男","日本","","","不通過","年齢","未設定","","不採用",""],
 [5,"バイトル","NRT",46232,"町田 侑里夏",21,"女","日本","","","通過","","面接調整中","","選考中",""],
 [6,"バイトル","NRT",46233,"三木 祐樹",38,"男","日本","","","不通過","年齢","未設定","","不採用",""],
 [7,"バイトル","NRT",46233,"金澤 拓也",19,"男","日本","","","通過","","面接調整中","","選考中",""],
 [8,"バイトル","NRT",46233,"NGUYEN NGOC DINH",25,"男","外国籍","","","不通過","日本語力","未設定","","不採用","N3"],
 [9,"バイトル","NRT",46236,"Hashini Navodya",21,"女","外国籍","","","確認中","","未設定","","選考中",""],
 [10,"バイトル","NRT",46236,"JARGALBAYAR ENKHBAYAN",21,"男","外国籍","","","確認中","","未設定","","選考中",""],
 [11,"バイトル","NRT",46237,"坂本 隼斗",24,"男","日本","","","確認中","","未設定","","選考中",""],
 [12,"バイトル","CTS",46236,"Hashini Navodya",21,"女","外国籍","","","確認中","","未設定","","選考中",""],
 [13,"マイナビバイト","CTS",46233,"竹内 愛",33,"女","日本","","","確認中","","未設定","","選考中",""],
 [14,"マイナビバイト","CTS",46233,"高砂 琴音",18,"女","日本","未経験者","学生","確認中","","未設定","","選考中","高校生（外航GATE経験者の可能性あり）"],
]
app.merge_cells("A1:P1")
t=app["A1"]; t.value="応募者データ（明細）　― 毎日ここを更新。ダッシュボード・社長報告は自動集計されます ―"
t.font=font(12,True,WHITE); t.fill=fill(NAVY); t.alignment=LEF; app.row_dimensions[1].height=24
for c,(h,w) in enumerate(zip(cols,widths),1):
    cell=app.cell(row=2,column=c,value=h); cell.font=font(10,True,WHITE); cell.fill=fill(BLUE)
    cell.alignment=CEN; cell.border=BORDER; app.column_dimensions[get_column_letter(c)].width=w
for i,row in enumerate(data):
    r=3+i
    for c,v in enumerate(row,1):
        cell=app.cell(row=r,column=c,value=(v if v!="" else None)); cell.font=font(10); cell.border=BORDER
        cell.alignment=CEN if c in(1,2,3,4,6,7,8,9,10,11,13,15) else LEF
        if c==4 and v!="": cell.number_format="yyyy/mm/dd"
app.freeze_panes="A3"
dv_targets={"媒体":"媒体","基地":"基地","性別":"性別","国籍":"国籍","経験":"経験","属性":"属性",
            "書類選考":"書類選考","書類NG理由":"書類NG","面接":"面接","面接NG理由":"面接NG","採用":"採用"}
for mkey,ckey in dv_targets.items():
    l,end=mrange[mkey]
    dv=DataValidation(type="list",formula1=f"=マスタ!${l}$3:${l}${end}",allow_blank=True)
    app.add_data_validation(dv); dv.add(f"{COL[ckey]}3:{COL[ckey]}{DATA_END}")

# =====================================================================
#  DAILY REPORT LOG
# =====================================================================
dl=wb.create_sheet("日次レポート")
dcols=["日付","応募数","書類通過","面接設定","面接実施","内定","承諾","担当者","備考"]
dwidths=[13,9,9,9,9,7,7,12,30]
dl.merge_cells("A1:I1"); t=dl["A1"]; t.value="日次レポート（毎日 EOD に担当者が1行追記）"
t.font=font(12,True,WHITE); t.fill=fill(NAVY); t.alignment=LEF; dl.row_dimensions[1].height=24
dl.merge_cells("A2:I2"); n=dl["A2"]
n.value=("使い方：応募者データを更新後、当日の累計値を下表に1行追記します（右の「現在の累計値」を転記でOK）。"
         "ダッシュボードの前日比・推移グラフに自動反映されます。灰色の行はサンプルです。実データ入力時に削除してください。")
n.font=font(9,False,GREY_H,it=True); n.alignment=LEF; dl.row_dimensions[2].height=30
for c,(h,w) in enumerate(zip(dcols,dwidths),1):
    cell=dl.cell(row=3,column=c,value=h); cell.font=font(10,True,WHITE); cell.fill=fill(BLUE)
    cell.alignment=CEN; cell.border=BORDER; dl.column_dimensions[get_column_letter(c)].width=w
samples=[[46234,12,1,1,0,0,0,"（サンプル）","← サンプル：削除して実データを入力"],
         [46235,13,2,2,0,0,0,"（サンプル）",""],
         [46236,14,2,2,0,0,0,"（サンプル）",""]]
DL_START=4
for i,row in enumerate(samples):
    r=DL_START+i
    for c,v in enumerate(row,1):
        cell=dl.cell(row=r,column=c,value=v); cell.font=font(10,color=GREY_H,it=True); cell.border=BORDER
        cell.alignment=CEN if c<=7 else LEF; cell.fill=fill(GREY_L)
        if c==1: cell.number_format="yyyy/mm/dd"
for r in range(DL_START+len(samples),DL_START+len(samples)+20):
    for c in range(1,10):
        cell=dl.cell(row=r,column=c); cell.border=BORDER
        if c==1: cell.number_format="yyyy/mm/dd"
dl.freeze_panes="A4"
dl["K3"]="現在の累計値（自動）"; dl["K3"].font=font(10,True,WHITE); dl["K3"].fill=fill("548235")
dl["K3"].alignment=CEN; dl.merge_cells("K3:L3")
dl.column_dimensions["K"].width=14; dl.column_dimensions["L"].width=9
live_funnel=[
 ("応募数",  f'=COUNTA({rng("氏名")})'),
 ("書類通過", f'={countif("書類選考","通過")}'),
 ("面接設定", f'={countif("面接","面接調整中")}+{countif("面接","面接日確定")}+{countif("面接","面接実施済み")}'),
 ("面接実施", f'={countif("面接","面接実施済み")}'),
 ("内定",    f'={countif("採用","採用")}+{countif("採用","入社予定")}+{countif("採用","採用辞退")}+{countif("採用","入社辞退")}'),
 ("承諾",    f'={countif("採用","採用")}+{countif("採用","入社予定")}'),
]
for i,(lab,fm) in enumerate(live_funnel):
    r=4+i
    c1=dl.cell(row=r,column=11,value=lab); c1.font=font(10,True); c1.border=BORDER; c1.alignment=LEF; c1.fill=fill(BLUE_L)
    c2=dl.cell(row=r,column=12,value=fm); c2.font=font(10); c2.border=BORDER; c2.alignment=CEN

# =====================================================================
#  KPI DASHBOARD
# =====================================================================
dash=wb.create_sheet("KPIダッシュボード")
dash.sheet_view.showGridLines=False
dash.column_dimensions["A"].width=2
for col,w in zip("BCDEFGH",[18,11,12,10,11,11,11]): dash.column_dimensions[col].width=w
dash.column_dimensions["I"].width=2
dash.merge_cells("B1:H1"); t=dash["B1"]; t.value="📊 短時間勤務者　採用KPIダッシュボード"
t.font=font(18,True,WHITE); t.fill=fill(NAVY); t.alignment=Alignment(horizontal="left",vertical="center")
dash.row_dimensions[1].height=34
dash.merge_cells("B2:H2"); st=dash["B2"]
st.value='=TEXT(MAX(日次レポート!$A$4:$A$'+str(DL_LAST)+'),"yyyy/mm/dd")&" 時点　｜　成田(NRT)・千歳(CTS) パート採用進捗　｜　数値は応募者データより自動集計"'
st.font=font(10,color=WHITE); st.fill=fill(BLUE); st.alignment=Alignment(horizontal="left",vertical="center")
dash.row_dimensions[2].height=20

def sec(ws,row,text,span="B{0}:H{0}"):
    ws.merge_cells(span.format(row)); c=ws.cell(row=row,column=2,value=text)
    c.font=font(12,True,WHITE); c.fill=fill(BLUE); c.alignment=LEF; ws.row_dimensions[row].height=22
def hdr(ws,row,labels,start=2):
    for i,lab in enumerate(labels):
        c=ws.cell(row=row,column=start+i,value=lab); c.font=font(10,True,WHITE); c.fill=fill(GREY_H)
        c.alignment=CEN; c.border=BORDER
def cf_pct(ws,rngstr):
    ws.conditional_formatting.add(rngstr,CellIsRule(operator="greaterThanOrEqual",formula=["1"],fill=fill(GREEN_F),font=Font(name=JP,color=GREEN_T,bold=True)))
    ws.conditional_formatting.add(rngstr,CellIsRule(operator="between",formula=["0.7","0.9999"],fill=fill(AMBER_F),font=Font(name=JP,color=AMBER_T)))
    ws.conditional_formatting.add(rngstr,CellIsRule(operator="lessThan",formula=["0.7"],fill=fill(RED_F),font=Font(name=JP,color=RED_T)))

# ---- target-setting block ----
def _inp(ws,r,c,val,fmt=None):
    cell=ws.cell(row=r,column=c,value=val); cell.font=font(11,True,"0000FF"); cell.fill=fill(INPUT)
    cell.alignment=CEN; cell.border=BORDER
    if fmt: cell.number_format=fmt
    return cell
def _lab(ws,r,c,val,bold=True,sz=10):
    cell=ws.cell(row=r,column=c,value=val); cell.font=font(sz,bold); cell.alignment=CEN; cell.border=BORDER
    return cell
sec(dash,4,"◎  目標設定 ─ 黄色セルに入力（採用目標から各段階の目標を自動逆算）")
c=dash.cell(row=5,column=2,value="採用目標（承諾・人）"); c.font=font(11,True); c.alignment=LEF; c.border=BORDER
_lab(dash,5,3,"NRT(成田)"); _inp(dash,5,4,10); _lab(dash,5,5,"CTS(千歳)"); _inp(dash,5,6,10); _lab(dash,5,7,"合計")
hc=dash.cell(row=5,column=8,value="=D5+F5"); hc.font=font(11,True); hc.fill=fill(BLUE_L); hc.alignment=CEN; hc.border=BORDER
c=dash.cell(row=6,column=2,value="歩留まり前提"); c.font=font(10,True); c.alignment=LEF; c.border=BORDER
for cc,tt in zip(range(3,8),["書類通過率","面接設定率","面接実施率","内定率","承諾率"]): _lab(dash,6,cc,tt)
c=dash.cell(row=7,column=2,value="（応募→承諾）"); c.font=font(9,color=GREY_H); c.alignment=LEF; c.border=BORDER
for cc,vv in zip(range(3,8),[0.5,0.7,0.85,0.6,0.8]): _inp(dash,7,cc,vv,"0%")

# ---- Section 1: funnel ----
sec(dash,9,"①  採用ファネル：目標(逆算) vs 実績（累計）")
hdr(dash,10,["指標","目標(逆算)","実績(累計)","達成率","前日","前日比"])
funnel_rows=["応募数","書類通過","面接設定","面接実施","内定","承諾"]
live=[fm for _,fm in live_funnel]
dlcol=["B","C","D","E","F","G"]
yieldref=["$C$7","$D$7","$E$7","$F$7","$G$7"]
F1=11
for i,name in enumerate(funnel_rows):
    r=F1+i
    dash.cell(row=r,column=2,value=name).font=font(11,True); dash.cell(row=r,column=2).alignment=LEF
    tfm="=$H$5" if i==5 else f"=ROUNDUP(C{r+1}/{yieldref[i]},0)"
    dash.cell(row=r,column=3,value=tfm).font=font(11,True); dash.cell(row=r,column=3).alignment=CEN
    dash.cell(row=r,column=4,value=live[i]).font=font(11,True); dash.cell(row=r,column=4).alignment=CEN
    dash.cell(row=r,column=5,value=f'=IFERROR(D{r}/C{r},"")').number_format="0.0%"; dash.cell(row=r,column=5).alignment=CEN
    dash.cell(row=r,column=6,value=f'=IFERROR(LOOKUP(1E+100,日次レポート!{dlcol[i]}$4:{dlcol[i]}${DL_LAST}),0)').alignment=CEN
    dc=dash.cell(row=r,column=7,value=f'=D{r}-F{r}'); dc.alignment=CEN; dc.number_format='+#,##0;-#,##0;0'
    for c in range(2,8):
        dash.cell(row=r,column=c).border=BORDER
        if c in(4,6): dash.cell(row=r,column=c).font=font(11)
cf_pct(dash,f"E{F1}:E{F1+5}")
dash.conditional_formatting.add(f"G{F1}:G{F1+5}",CellIsRule(operator="greaterThan",formula=["0"],fill=fill(GREEN_F),font=Font(name=JP,color=GREEN_T,bold=True)))
dash.conditional_formatting.add(f"G{F1}:G{F1+5}",CellIsRule(operator="lessThan",formula=["0"],fill=fill(RED_F),font=Font(name=JP,color=RED_T,bold=True)))

# ---- Section 2: conversion ----
CV=18
sec(dash,CV,"②  転換率（歩留まり）")
hdr(dash,CV+1,["指標","値","計算式・定義"])
conv=[("書類通過率",f"=IFERROR(D{F1+1}/D{F1},\"\")","書類通過 ÷ 応募数"),
      ("面接設定率",f"=IFERROR(D{F1+2}/D{F1+1},\"\")","面接設定 ÷ 書類通過"),
      ("面接実施率",f"=IFERROR(D{F1+3}/D{F1+2},\"\")","面接実施 ÷ 面接設定"),
      ("内定率",f"=IFERROR(D{F1+4}/D{F1+3},\"\")","内定 ÷ 面接実施"),
      ("承諾率",f"=IFERROR(D{F1+5}/D{F1+4},\"\")","承諾 ÷ 内定")]
for i,(name,fm,note) in enumerate(conv):
    r=CV+2+i
    dash.cell(row=r,column=2,value=name).font=font(11,True); dash.cell(row=r,column=2).alignment=LEF
    vc=dash.cell(row=r,column=3,value=fm); vc.number_format="0.0%"; vc.alignment=CEN; vc.font=font(11,True,BLUE)
    dash.merge_cells(start_row=r,start_column=4,end_row=r,end_column=7)
    nc=dash.cell(row=r,column=4,value=note); nc.font=font(9,color=GREY_H); nc.alignment=LEF
    for c in range(2,8): dash.cell(row=r,column=c).border=BORDER

# ---- Section 3: by media ----
MB=26
sec(dash,MB,"③  媒体別 実績")
hdr(dash,MB+1,["媒体","応募数","書類通過","面接設定","面接実施","内定","承諾"])
media=master["媒体"]
def seg_formulas(segkey,segval):
    base=[(segkey,segval)]
    return [
      f'=COUNTIFS({rng(segkey)},"{segval}")',
      f'=COUNTIFS({rng(segkey)},"{segval}",{rng("書類選考")},"通過")',
      "="+ "+".join(countifs(base+[("面接",x)]) for x in ["面接調整中","面接日確定","面接実施済み"]),
      f'=COUNTIFS({rng(segkey)},"{segval}",{rng("面接")},"面接実施済み")',
      "="+ "+".join(countifs(base+[("採用",x)]) for x in ["採用","入社予定","採用辞退","入社辞退"]),
      "="+ "+".join(countifs(base+[("採用",x)]) for x in ["採用","入社予定"]),
    ]
for i,m in enumerate(media):
    r=MB+2+i
    dash.cell(row=r,column=2,value=m).font=font(10,True); dash.cell(row=r,column=2).alignment=LEF
    for j,fm in enumerate(seg_formulas("媒体",m)):
        dash.cell(row=r,column=3+j,value=fm).alignment=CEN; dash.cell(row=r,column=3+j).font=font(10)
    for c in range(2,9): dash.cell(row=r,column=c).border=BORDER
    if i%2==1:
        for c in range(2,9): dash.cell(row=r,column=c).fill=fill(GREY_L)
tr=MB+2+len(media)
dash.cell(row=tr,column=2,value="合計").font=font(10,True,WHITE); dash.cell(row=tr,column=2).fill=fill(BLUE); dash.cell(row=tr,column=2).alignment=LEF
for j in range(6):
    col=get_column_letter(3+j)
    cc=dash.cell(row=tr,column=3+j,value=f'=SUM({col}{MB+2}:{col}{MB+1+len(media)})')
    cc.font=font(10,True,WHITE); cc.fill=fill(BLUE); cc.alignment=CEN
for c in range(2,9): dash.cell(row=tr,column=c).border=BORDER

# ---- Section 4: by base ----
BB=tr+2
sec(dash,BB,"④  基地別 実績（NRT=成田 / CTS=千歳）")
hdr(dash,BB+1,["基地","応募数","書類通過","面接設定","面接実施","内定","承諾"])
bases=[("NRT","成田"),("CTS","千歳")]
for i,(code,label) in enumerate(bases):
    r=BB+2+i
    dash.cell(row=r,column=2,value=f"{code}（{label}）").font=font(10,True); dash.cell(row=r,column=2).alignment=LEF
    for j,fm in enumerate(seg_formulas("基地",code)):
        dash.cell(row=r,column=3+j,value=fm).alignment=CEN; dash.cell(row=r,column=3+j).font=font(10)
    for c in range(2,9): dash.cell(row=r,column=c).border=BORDER

# ---- Section 5: per-base hire goal ----
PT=BB+2+len(bases)+1
sec(dash,PT,"⑤  拠点別 採用目標 達成状況（承諾＝採用人数）")
hdr(dash,PT+1,["拠点","採用目標","承諾実績","達成率"])
pt_rows=[("NRT（成田）","$D$5",f"H{BB+2}"),("CTS（千歳）","$F$5",f"H{BB+3}")]
for i,(lab,goal,act) in enumerate(pt_rows):
    r=PT+2+i
    dash.cell(row=r,column=2,value=lab).font=font(10,True); dash.cell(row=r,column=2).alignment=LEF
    dash.cell(row=r,column=3,value=f"={goal}").alignment=CEN; dash.cell(row=r,column=3).font=font(10)
    dash.cell(row=r,column=4,value=f"={act}").alignment=CEN; dash.cell(row=r,column=4).font=font(10)
    rc=dash.cell(row=r,column=5,value=f'=IFERROR({act}/{goal},"")'); rc.number_format="0.0%"; rc.alignment=CEN
    for c in range(2,6): dash.cell(row=r,column=c).border=BORDER
r=PT+2+len(pt_rows)
dash.cell(row=r,column=2,value="合計").font=font(10,True,WHITE); dash.cell(row=r,column=2).fill=fill(BLUE); dash.cell(row=r,column=2).alignment=LEF
dash.cell(row=r,column=3,value="=$H$5").font=font(10,True,WHITE); dash.cell(row=r,column=3).fill=fill(BLUE); dash.cell(row=r,column=3).alignment=CEN
dash.cell(row=r,column=4,value=f"=H{BB+2}+H{BB+3}").font=font(10,True,WHITE); dash.cell(row=r,column=4).fill=fill(BLUE); dash.cell(row=r,column=4).alignment=CEN
tc=dash.cell(row=r,column=5,value=f'=IFERROR((H{BB+2}+H{BB+3})/$H$5,"")'); tc.number_format="0.0%"; tc.font=font(10,True,WHITE); tc.fill=fill(BLUE); tc.alignment=CEN
for c in range(2,6): dash.cell(row=r,column=c).border=BORDER
cf_pct(dash,f"E{PT+2}:E{PT+2+len(pt_rows)}")

# ---- Section 6: applicant attributes (累計) ----
AT=PT+2+len(pt_rows)+2
sec(dash,AT,"⑥  応募者属性（累計）")
def attr_block(top, title, items, total_ref=None):
    """items: list of (label, formula). renders 指標|人数|構成比 table starting at row top."""
    hdr(dash,top,[title,"人数","構成比"])
    n=len(items)
    for i,(lab,fm) in enumerate(items):
        r=top+1+i
        dash.cell(row=r,column=2,value=lab).font=font(10); dash.cell(row=r,column=2).alignment=LEF
        vc=dash.cell(row=r,column=3,value=fm); vc.alignment=CEN; vc.font=font(10)
        pc=dash.cell(row=r,column=4,value=f'=IFERROR(C{r}/SUM(C{top+1}:C{top+n}),"")'); pc.number_format="0.0%"; pc.alignment=CEN
        for c in range(2,5): dash.cell(row=r,column=c).border=BORDER
    return top+1+n
OUBO=f'COUNTA({rng("氏名")})'
# 男女
r_after=attr_block(AT+1,"男女",[("男",f'={countif("性別","男")}'),("女",f'={countif("性別","女")}'),
    ("未入力",f'={OUBO}-{countif("性別","男")}-{countif("性別","女")}')])
# 国籍
r_after=attr_block(r_after+1,"国籍",[("日本",f'={countif("国籍","日本")}'),("外国籍",f'={countif("国籍","外国籍")}'),
    ("未入力",f'={OUBO}-{countif("国籍","日本")}-{countif("国籍","外国籍")}')])
# 経験
r_after=attr_block(r_after+1,"経験",[("経験者",f'={countif("経験","経験者")}'),("未経験者",f'={countif("経験","未経験者")}'),
    ("未入力",f'={OUBO}-{countif("経験","経験者")}-{countif("経験","未経験者")}')])
# 就業区分（属性）
attr_items=[(x,f'={countif("属性",x)}') for x in ["学生","フリーター","社会人","主婦・主夫","その他"]]
subtract="".join(f'-{countif("属性",x)}' for x in ["学生","フリーター","社会人","主婦・主夫","その他"])
attr_items.append(("未入力",f'={OUBO}{subtract}'))
r_after=attr_block(r_after+1,"就業区分",attr_items)
# 年齢分布
age_items=[("〜19歳",f'={countif("年齢","<=19")}'),
           ("20代",f'=COUNTIFS({rng("年齢")},">=20",{rng("年齢")},"<=29")'),
           ("30代",f'=COUNTIFS({rng("年齢")},">=30",{rng("年齢")},"<=39")'),
           ("40代",f'=COUNTIFS({rng("年齢")},">=40",{rng("年齢")},"<=49")'),
           ("50代",f'=COUNTIFS({rng("年齢")},">=50",{rng("年齢")},"<=59")'),
           ("60歳〜",f'={countif("年齢",">=60")}')]
AGE_top=r_after+1
r_after=attr_block(AGE_top,"年齢分布",age_items)

# ---- legend ----
LG=r_after+2
dash.merge_cells(f"B{LG}:H{LG}"); dash.cell(row=LG,column=2,value="◆ 凡例・定義").font=font(10,True,BLUE)
notes=[
 "・黄色セル＝入力。採用目標（NRT/CTSの承諾人数）と歩留まり前提を入れると、各段階の目標を自動逆算します。",
 "・目標列＝承諾目標から逆算（承諾20名→内定25→面接実施42→面接設定50→書類通過72→応募144）。歩留まりを変えれば再計算されます。",
 "・実績・属性は［応募者データ］から自動集計。担当者は毎日そこを更新するだけです。経験／就業区分は各明細で入力してください。",
 "・前日／前日比は［日次レポート］の最終記録行との差分。毎日EODに1行追記してください。",
 "・面接設定＝「面接調整中／面接日確定／面接実施済み」。内定＝採用が「採用／入社予定／採用辞退／入社辞退」。承諾＝「採用／入社予定」。",
]
for i,tx in enumerate(notes):
    r=LG+1+i; dash.merge_cells(f"B{r}:H{r}")
    dash.cell(row=r,column=2,value=tx).font=font(9,color=GREY_H); dash.cell(row=r,column=2).alignment=LEF

# ---- charts ----
ch1=BarChart(); ch1.type="bar"; ch1.title="採用ファネル（実績）"; ch1.height=7; ch1.width=12
ch1.add_data(Reference(dash,min_col=4,min_row=F1,max_row=F1+5)); ch1.set_categories(Reference(dash,min_col=2,min_row=F1,max_row=F1+5))
ch1.legend=None; ch1.dataLabels=DataLabelList(); ch1.dataLabels.showVal=True; dash.add_chart(ch1,"J4")
ch2=BarChart(); ch2.type="col"; ch2.title="媒体別 応募数"; ch2.height=7; ch2.width=12
ch2.add_data(Reference(dash,min_col=3,min_row=MB+1,max_row=MB+1+len(media)),titles_from_data=True)
ch2.set_categories(Reference(dash,min_col=2,min_row=MB+2,max_row=MB+1+len(media)))
ch2.legend=None; ch2.dataLabels=DataLabelList(); ch2.dataLabels.showVal=True; dash.add_chart(ch2,"J20")
ch3=LineChart(); ch3.title="推移（日次）"; ch3.height=7; ch3.width=12
ch3.add_data(Reference(dl,min_col=2,max_col=7,min_row=3,max_row=DL_START+len(samples)+9),titles_from_data=True)
ch3.set_categories(Reference(dl,min_col=1,min_row=DL_START,max_row=DL_START+len(samples)+9)); ch3.markers=True
dash.add_chart(ch3,"J36")
ch4=BarChart(); ch4.type="col"; ch4.title="年齢分布"; ch4.height=7; ch4.width=12
ch4.add_data(Reference(dash,min_col=3,min_row=AGE_top,max_row=AGE_top+len(age_items)))
ch4.set_categories(Reference(dash,min_col=2,min_row=AGE_top+1,max_row=AGE_top+len(age_items)))
ch4.legend=None; ch4.dataLabels=DataLabelList(); ch4.dataLabels.showVal=True; dash.add_chart(ch4,"J52")
dash.sheet_view.zoomScale=100; dash.sheet_properties.tabColor=NAVY

# =====================================================================
#  PRESIDENT DAILY REPORT (社長報告)
# =====================================================================
rep=wb.create_sheet("社長報告")
rep.sheet_view.showGridLines=False
rep.column_dimensions["A"].width=2
for col,w in zip("BCDEFG",[20,13,13,13,13,13]): rep.column_dimensions[col].width=w
D="KPIダッシュボード"
def rsec(row,text,span="B{0}:G{0}"):
    rep.merge_cells(span.format(row)); c=rep.cell(row=row,column=2,value=text)
    c.font=font(11,True,WHITE); c.fill=fill(BLUE); c.alignment=LEF; rep.row_dimensions[row].height=20
# Title
rep.merge_cells("B1:G1"); t=rep["B1"]; t.value="採用活動　日次報告書"
t.font=font(18,True,WHITE); t.fill=fill(NAVY); t.alignment=CEN; rep.row_dimensions[1].height=32
rep.merge_cells("B2:G2"); s=rep["B2"]
s.value='="報告日：" & TEXT(MAX(日次レポート!$A$4:$A$'+str(DL_LAST)+'),"yyyy年m月d日") & "　｜　短時間勤務者（成田NRT・千歳CTS）　｜　宛先：社長"'
s.font=font(10,color=INK); s.alignment=CEN; rep.row_dimensions[2].height=18
rep.merge_cells("B3:G3"); s=rep["B3"]; s.value="担当者：________________"
s.font=font(9,color=GREY_H); s.alignment=Alignment(horizontal="right",vertical="center")

# Highlight cards row (承諾 / 応募 / 面接実施)
def card(col,label,val_ref,sub_ref,daily_ref):
    rep.merge_cells(start_row=5,start_column=col,end_row=5,end_column=col+1)
    lc=rep.cell(row=5,column=col,value=label); lc.font=font(10,True,WHITE); lc.fill=fill(NAVY); lc.alignment=CEN
    rep.merge_cells(start_row=6,start_column=col,end_row=7,end_column=col+1)
    vc=rep.cell(row=6,column=col,value=val_ref); vc.font=font(22,True,NAVY); vc.alignment=CEN; vc.fill=fill(BLUE_L)
    rep.merge_cells(start_row=8,start_column=col,end_row=8,end_column=col+1)
    sc=rep.cell(row=8,column=col,value=sub_ref); sc.font=font(9,color=GREY_H); sc.alignment=CEN; sc.fill=fill(BLUE_L)
    for rr in (5,6,8):
        for cc in (col,col+1): rep.cell(row=rr,column=cc).border=BORDER
card(2,"承諾（採用）", f'={D}!D{F1+5}&" / "&{D}!C{F1+5}&" 名"', f'=" 達成率 "&TEXT({D}!E{F1+5},"0.0%")&"　前日比 "&TEXT({D}!G{F1+5},"+0;-0;0")', None)
card(4,"応募数（累計）", f'={D}!D{F1}&" 名"', f'=" 目標 "&{D}!C{F1}&"　達成率 "&TEXT({D}!E{F1},"0.0%")&"　前日比 "&TEXT({D}!G{F1},"+0;-0;0")', None)
card(6,"面接実施（累計）", f'={D}!D{F1+3}&" 名"', f'=" 目標 "&{D}!C{F1+3}&"　前日比 "&TEXT({D}!G{F1+3},"+0;-0;0")', None)
rep.row_dimensions[6].height=26

# Funnel summary
rsec(10,"■ 採用ファネル（目標 vs 実績）")
hdr(rep,11,["指標","目標","実績","達成率","前日比"])
for i,name in enumerate(funnel_rows):
    r=12+i; dr=F1+i
    rep.cell(row=r,column=2,value=name).font=font(10,True); rep.cell(row=r,column=2).alignment=LEF
    rep.cell(row=r,column=3,value=f"={D}!C{dr}").alignment=CEN; rep.cell(row=r,column=3).font=font(10)
    rep.cell(row=r,column=4,value=f"={D}!D{dr}").font=font(10,True); rep.cell(row=r,column=4).alignment=CEN
    ac=rep.cell(row=r,column=5,value=f"={D}!E{dr}"); ac.number_format="0.0%"; ac.alignment=CEN; ac.font=font(10)
    gc=rep.cell(row=r,column=6,value=f"={D}!G{dr}"); gc.number_format='+#,##0;-#,##0;0'; gc.alignment=CEN
    for c in range(2,7): rep.cell(row=r,column=c).border=BORDER
cf_pct(rep,"E12:E17")
rep.conditional_formatting.add("F12:F17",CellIsRule(operator="greaterThan",formula=["0"],fill=fill(GREEN_F),font=Font(name=JP,color=GREEN_T,bold=True)))
rep.conditional_formatting.add("F12:F17",CellIsRule(operator="lessThan",formula=["0"],fill=fill(RED_F),font=Font(name=JP,color=RED_T,bold=True)))

# Per-base hire
rsec(19,"■ 拠点別 採用目標 達成（承諾）")
hdr(rep,20,["拠点","採用目標","承諾実績","達成率"],start=2)
rep.merge_cells("E20:G20")
for i,(lab,drow) in enumerate([("NRT（成田）",PT+2),("CTS（千歳）",PT+3),("合計",PT+4)]):
    r=21+i
    rep.cell(row=r,column=2,value=lab).font=font(10,True); rep.cell(row=r,column=2).alignment=LEF
    rep.cell(row=r,column=3,value=f"={D}!C{drow}").alignment=CEN; rep.cell(row=r,column=3).font=font(10)
    rep.cell(row=r,column=4,value=f"={D}!D{drow}").alignment=CEN; rep.cell(row=r,column=4).font=font(10,True)
    rep.merge_cells(start_row=r,start_column=5,end_row=r,end_column=7)
    rc=rep.cell(row=r,column=5,value=f"={D}!E{drow}"); rc.number_format="0.0%"; rc.alignment=CEN; rc.font=font(10)
    for c in range(2,8): rep.cell(row=r,column=c).border=BORDER

# Attributes (累計) — two mini tables side by side
rsec(25,"■ 応募者属性（累計）")
def mini(top,left,title,items):
    """render title + rows(label|人数|比) in columns left..left+2"""
    for j,txt in enumerate([title,"人数","比率"]):
        cc=rep.cell(row=top,column=left+j,value=txt); cc.font=font(9,True,WHITE); cc.fill=fill(GREY_H); cc.alignment=CEN; cc.border=BORDER
    n=len(items)
    for i,(lab,fm) in enumerate(items):
        r=top+1+i
        rep.cell(row=r,column=left,value=lab).font=font(9); rep.cell(row=r,column=left).alignment=LEF
        vc=rep.cell(row=r,column=left+1,value=fm); vc.alignment=CEN; vc.font=font(9)
        col1=get_column_letter(left+1)
        pc=rep.cell(row=r,column=left+2,value=f'=IFERROR({col1}{r}/SUM({col1}{top+1}:{col1}{top+n}),"")'); pc.number_format="0%"; pc.alignment=CEN; pc.font=font(9)
        for c in range(left,left+3): rep.cell(row=r,column=c).border=BORDER
# left column tables: 男女, 経験  |  right column: 国籍, 就業区分  |  下: 年齢
gender=[("男",f'={countif("性別","男")}'),("女",f'={countif("性別","女")}')]
exp=[("経験者",f'={countif("経験","経験者")}'),("未経験者",f'={countif("経験","未経験者")}'),("未入力",f'={OUBO}-{countif("経験","経験者")}-{countif("経験","未経験者")}')]
nat=[("日本",f'={countif("国籍","日本")}'),("外国籍",f'={countif("国籍","外国籍")}')]
occ=[(x,f'={countif("属性",x)}') for x in ["学生","フリーター","社会人","主婦・主夫","その他"]]
occ.append(("未入力",f'={OUBO}{subtract}'))
mini(26,2,"男女",gender)
mini(30,2,"経験",exp)
mini(26,5,"国籍",nat)
mini(30,5,"就業区分",occ)
# 年齢 across bottom (label row + counts)
AGEr=38
rep.cell(row=AGEr,column=2,value="年齢分布").font=font(9,True,WHITE); rep.cell(row=AGEr,column=2).fill=fill(GREY_H); rep.cell(row=AGEr,column=2).alignment=CEN; rep.cell(row=AGEr,column=2).border=BORDER
age_labels=["〜19","20代","30代","40代","50代","60〜"]
age_fms=[f'={countif("年齢","<=19")}',
         f'=COUNTIFS({rng("年齢")},">=20",{rng("年齢")},"<=29")',
         f'=COUNTIFS({rng("年齢")},">=30",{rng("年齢")},"<=39")',
         f'=COUNTIFS({rng("年齢")},">=40",{rng("年齢")},"<=49")',
         f'=COUNTIFS({rng("年齢")},">=50",{rng("年齢")},"<=59")',
         f'={countif("年齢",">=60")}']
for j in range(len(age_labels)):
    hc=rep.cell(row=AGEr,column=3+j,value=age_labels[j]); hc.font=font(9,True); hc.alignment=CEN; hc.border=BORDER; hc.fill=fill(BLUE_L)
    vc=rep.cell(row=AGEr+1,column=3+j,value=age_fms[j]); vc.font=font(10,True); vc.alignment=CEN; vc.border=BORDER
lc=rep.cell(row=AGEr+1,column=2,value="人数"); lc.font=font(9); lc.alignment=CEN; lc.border=BORDER

# Comment box
rsec(41,"■ 特記事項・本日の動き（担当者記入）")
rep.merge_cells("B42:G45")
cb=rep.cell(row=42,column=2,value="（例）本日バイトル千歳より新規3名応募。町田様 面接日程調整中。年齢超過による書類NGが継続、媒体訴求を検討。")
cb.font=font(10,color="0000FF"); cb.fill=fill(INPUT); cb.alignment=Alignment(horizontal="left",vertical="top",wrap_text=True)
for r in range(42,46):
    for c in range(2,8): rep.cell(row=r,column=c).border=BORDER
# print setup
rep.print_area="A1:G46"
rep.page_setup.orientation="portrait"; rep.page_setup.fitToWidth=1; rep.page_setup.fitToHeight=1
rep.sheet_properties.pageSetUpPr=openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
rep.page_margins=openpyxl.worksheet.page.PageMargins(left=0.4,right=0.4,top=0.5,bottom=0.4)
rep.sheet_properties.tabColor="C00000"

# ---- sheet order: 社長報告, KPIダッシュボード, 日次レポート, 応募者データ, マスタ ----
order=["社長報告","KPIダッシュボード","日次レポート","応募者データ","マスタ"]
wb._sheets.sort(key=lambda ws: order.index(ws.title))
wb.active=0
wb.calculation=CalcProperties(calcId=0,fullCalcOnLoad=True)
wb.save(OUT)
print("saved",OUT); print("sheets:",wb.sheetnames)
